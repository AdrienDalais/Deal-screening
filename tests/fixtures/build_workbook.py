"""Generate a financials workbook from a CompanyTruth.

Layout: a header block stating currency and units, a merged header
separating historicals from forecasts, one column per period, one row
per line item, and a checks block.

Costs are written negative, matching the convention real working models
use, so that gross profit is revenue *plus* cost of sales.

Defects are applied here, explicitly and by key, so that every distortion
in the output can be traced to one line of code.
"""

from __future__ import annotations

import re
import zipfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from storage import enums as E
from tests.fixtures.truth import CompanyTruth

FIRST_PERIOD_COL = 4          # column D
PERIOD_HEADER_ROW = 8
LABEL_COL = 2                 # column B
# openpyxl stamps the workbook with the time it was saved, so a
# rebuild differs from its predecessor by a few bytes. Pinned, so
# that a difference in a fixture means a difference that matters.
FIXTURE_TIMESTAMP = datetime(2026, 1, 1, 0, 0, 0)

BOLD = Font(bold=True)


def _normalise_archive(path: Path) -> None:
    """Rewrite an .xlsx so it is byte-identical on every rebuild.

    An .xlsx is a zip, and a zip records the time each entry was written.
    Pinning the workbook's own created/modified properties is not enough:
    the archive itself still differs on every save, so a rebuild looks
    like a change when nothing has changed.

    Entries are rewritten with a fixed timestamp, in sorted order, with
    fixed compression. The spreadsheet is unaffected — only the envelope
    around it.
    """
    with zipfile.ZipFile(path) as archive:
        entries = sorted(archive.namelist())
        contents = {name: archive.read(name) for name in entries}

    # openpyxl rewrites dcterms:modified at save time whatever the
    # workbook properties say, so it is pinned here instead.
    core = "docProps/core.xml"
    if core in contents:
        contents[core] = re.sub(
            rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
            rb"\g<1>2026-01-01T00:00:00Z\g<2>", contents[core])

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED,
                         compresslevel=9) as archive:
        for name in entries:
            info = zipfile.ZipInfo(name, date_time=(2026, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            archive.writestr(info, contents[name])
UNDERLINE = Border(bottom=Side(style="thin"))


def _scaled(value: float, unit: E.Unit) -> float:
    """Whole units -> the workbook's scale."""
    return round(value / unit.multiplier, 4)


def build(truth: CompanyTruth, out_dir: Path) -> tuple[Path, dict[str, int]]:
    """Returns the path and a metric -> row-number map.

    The row map lets the answer key cite the exact cell each figure was
    written to. Hardcoding cell references in the key would break the
    moment a line item was added above them.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Scorecard"

    # --- header block ------------------------------------------------
    # Units are a required header field. A sheet whose scale is only
    # implied is ambiguous by a factor of 1,000 or more, so the
    # generator cannot write one without stating it.
    header = [
        ("Company", truth.name),
        ("Currency", truth.currency),
        ("Units", truth.workbook_unit.value),
        ("Year end (EoP)", ["", "Jan", "Feb", "Mar", "Apr", "May", "Jun",
                            "Jul", "Aug", "Sep", "Oct", "Nov",
                            "Dec"][truth.fiscal_year_end_month]),
    ]
    for offset, (label, value) in enumerate(header):
        ws.cell(row=2 + offset, column=LABEL_COL, value=label).font = BOLD
        ws.cell(row=2 + offset, column=LABEL_COL + 1, value=value)

    # --- period headers ----------------------------------------------
    # A merged cell spanning the actual columns, and another spanning the
    # forecasts. This is how real models signal basis: not with a suffix
    # on every column, but with a heading several rows above.
    header_row, period_row = 7, 8
    actual_cols, forecast_cols = [], []
    for index, period in enumerate(truth.periods):
        col = FIRST_PERIOD_COL + index
        suffix = "A" if period.basis == E.Basis.ACTUAL else "E"
        cell = ws.cell(row=period_row, column=col, value=f"{period.label}{suffix}")
        cell.font = BOLD
        cell.alignment = Alignment(horizontal="right")
        cell.border = UNDERLINE
        (actual_cols if period.basis == E.Basis.ACTUAL else forecast_cols).append(col)

    def _span(cols: list[int], text: str) -> None:
        if not cols:
            return
        first, last = min(cols), max(cols)
        ws.merge_cells(start_row=header_row, start_column=first,
                       end_row=header_row, end_column=last)
        cell = ws.cell(row=header_row, column=first, value=text)
        cell.font = BOLD
        cell.alignment = Alignment(horizontal="center")

    # ---- DEFECT: forecast_presented_as_actual -----------------------
    # The merged "Historicals" header covers one column more than it
    # should, placing the first forecast period inside the historical
    # block. The column's own label still carries the forecast suffix, so
    # two signals within one sheet disagree — which is how this survives
    # review in practice.
    if any(d.kind == "forecast_presented_as_actual" for d in truth.defects) \
            and forecast_cols:
        actual_cols = actual_cols + [forecast_cols[0]]
        forecast_cols = forecast_cols[1:]

    _span(actual_cols, "Historicals")
    _span(forecast_cols, "Management forecast")

    # --- line items ---------------------------------------------------
    row = period_row + 2
    row_of_metric: dict[str, int] = {}

    for item in truth.line_items:
        ws.cell(row=row, column=LABEL_COL, value=item.workbook_label)
        for index, period in enumerate(truth.periods):
            raw = item.values.get(period.label)
            if raw is None:
                continue
            value = _scaled(raw, truth.workbook_unit)
            if item.is_cost:
                value = -value

            # ---- DEFECT: thornbury_rekeying_scale_slip --------------
            # FY23 cost of sales re-keyed at the CIM's scale (millions)
            # into a sheet denominated in thousands.
            if (any(d.key == "thornbury_rekeying_scale_slip" for d in truth.defects)
                    and item.metric == "cogs" and period.label == "FY23"):
                value = value / 1000.0

            ws.cell(row=row, column=FIRST_PERIOD_COL + index, value=value)
        row_of_metric[item.metric] = row
        row += 1

    # --- checks -------------------------------------------------------
    # Every check is a live formula with a stated expectation. A check
    # that cannot be evaluated must say so in words, never render blank
    # or as an error code: all three read as 'fine' to a tired reader.
    row += 1
    ws.cell(row=row, column=LABEL_COL, value="Checks").font = BOLD
    row += 1

    checks = [
        ("Revenue less components (expect 0)", "revenue",
         ["recurring_revenue", "non_recurring_revenue"]),
        ("Gross profit less revenue and COGS (expect 0)", "gross_profit",
         ["revenue", "cogs"]),
        ("Total OPEX less components (expect 0)", "total_opex",
         ["personnel_expenses", "sales_and_marketing", "general_and_admin"]),
    ]
    for label, total_metric, parts in checks:
        if total_metric not in row_of_metric:
            continue
        if any(p not in row_of_metric for p in parts):
            ws.cell(row=row, column=LABEL_COL, value=label)
            ws.cell(row=row, column=FIRST_PERIOD_COL,
                    value="not evaluable: component row missing")
            row += 1
            continue
        ws.cell(row=row, column=LABEL_COL, value=label)
        for index in range(len(truth.periods)):
            col = get_column_letter(FIRST_PERIOD_COL + index)
            total_ref = f"{col}{row_of_metric[total_metric]}"
            part_refs = "+".join(f"{col}{row_of_metric[p]}" for p in parts)
            ws.cell(row=row, column=FIRST_PERIOD_COL + index,
                    value=f"={total_ref}-({part_refs})")
        row += 1

    # ---- DEFECT: check_cannot_be_evaluated --------------------------
    # A control that references a name which does not exist. It reports
    # an error permanently while being displayed as though it worked,
    # which is how a broken check comes to be read as decoration.
    if any(d.kind == "check_cannot_be_evaluated" for d in truth.defects):
        ws.cell(row=row, column=LABEL_COL,
                value="EBITDA less P&L EBITDA (expect 0)")
        for index in range(len(truth.periods)):
            ws.cell(row=row, column=FIRST_PERIOD_COL + index,
                    value="=Adj_EBITDA_PL-"
                          f"{get_column_letter(FIRST_PERIOD_COL + index)}"
                          f"{row_of_metric.get('adjusted_ebitda', 1)}")
        row += 1

    ws.column_dimensions["B"].width = 34
    for index in range(len(truth.periods)):
        ws.column_dimensions[get_column_letter(FIRST_PERIOD_COL + index)].width = 13

    wb.properties.created = FIXTURE_TIMESTAMP
    wb.properties.modified = FIXTURE_TIMESTAMP
    wb.properties.creator = "deal-screening-workbench"
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"{truth.slug}_financials.xlsx"
    wb.save(path)
    _normalise_archive(path)
    return path, row_of_metric


MANAGEMENT_ACCOUNT_ROWS = ["revenue", "recurring_revenue",
                           "non_recurring_revenue", "cogs", "gross_profit",
                           "total_opex", "adjusted_ebitda"]


def build_management_accounts(truth: CompanyTruth,
                              out_dir: Path) -> tuple[Path, dict[str, int]]:
    """A shorter, later document restating some of the same periods.

    Real deals produce these constantly: the CIM goes out, then accounts
    are issued that treat a divested line differently. Neither document
    is wrong. The system must hold both and let the analyst record which
    basis the paper is written on — which is why financial_observations
    has no uniqueness constraint on (deal, metric, period).
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Management Accounts"

    header = [
        ("Company", truth.name),
        ("Currency", truth.currency),
        ("Units", truth.workbook_unit.value),
        ("Year end (EoP)", ["", "Jan", "Feb", "Mar", "Apr", "May", "Jun",
                            "Jul", "Aug", "Sep", "Oct", "Nov",
                            "Dec"][truth.fiscal_year_end_month]),
        ("Basis", "Continuing operations"),
    ]
    for offset, (label, value) in enumerate(header):
        ws.cell(row=2 + offset, column=LABEL_COL, value=label).font = BOLD
        ws.cell(row=2 + offset, column=LABEL_COL + 1, value=value)

    actuals = [p for p in truth.periods if p.basis == E.Basis.ACTUAL]
    for index, period in enumerate(actuals):
        cell = ws.cell(row=PERIOD_HEADER_ROW, column=FIRST_PERIOD_COL + index,
                       value=f"{period.label}A")
        cell.font = BOLD
        cell.alignment = Alignment(horizontal="right")
        cell.border = UNDERLINE

    row = PERIOD_HEADER_ROW + 2
    row_of_metric: dict[str, int] = {}
    for metric in MANAGEMENT_ACCOUNT_ROWS:
        try:
            item = truth.item(metric)
        except KeyError:
            continue
        ws.cell(row=row, column=LABEL_COL, value=item.workbook_label)
        for index, period in enumerate(actuals):
            # A restated figure takes precedence over the CIM's.
            raw = truth.management_accounts.get(metric, {}).get(
                period.label, item.values.get(period.label))
            if raw is None:
                continue
            value = _scaled(raw, truth.workbook_unit)
            ws.cell(row=row, column=FIRST_PERIOD_COL + index,
                    value=-value if item.is_cost else value)
        row_of_metric[metric] = row
        row += 1

    ws.column_dimensions["B"].width = 34
    wb.properties.created = FIXTURE_TIMESTAMP
    wb.properties.modified = FIXTURE_TIMESTAMP
    wb.properties.creator = "deal-screening-workbench"
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"{truth.slug}_management_accounts.xlsx"
    wb.save(path)
    _normalise_archive(path)
    return path, row_of_metric
