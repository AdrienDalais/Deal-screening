"""Checks that must hold for every fixture company.

tests/test_fixtures.py covers Thornbury in depth. This file runs the
structural guarantees across all three, so a company added later cannot
quietly ship with a broken answer key.
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from tests.fixtures.build_all import COMPANIES
from tests.fixtures.truth import CompanyTruth

FIXTURES = Path(__file__).resolve().parent / "fixtures"
IDS = [c.slug for c in COMPANIES]


@pytest.fixture(params=COMPANIES, ids=IDS)
def company(request) -> CompanyTruth:
    return request.param


def _key(truth: CompanyTruth) -> dict:
    path = FIXTURES / truth.slug / "expected.json"
    if not path.exists():
        pytest.skip("fixtures not built; run python -m tests.fixtures.build_all")
    return json.loads(path.read_text())


# ---------------------------------------------------------------------
# Every company
# ---------------------------------------------------------------------

def test_the_truth_is_internally_consistent(company):
    """Before defects are applied, every total equals its components —
    on each breakdown axis independently."""
    assert company.check_internal_consistency() == []


def test_every_workbook_figure_sits_in_the_cell_the_key_cites(company):
    key = _key(company)
    sheet = load_workbook(FIXTURES / company.slug /
                          f"{company.slug}_financials.xlsx")["Scorecard"]
    doc = next(d for d in key["documents"]
               if d["document_type"] == "financials_workbook")
    for obs in doc["observations"]:
        actual = sheet[obs["source_cell"]].value
        assert actual == pytest.approx(obs["value_original"], abs=1e-6), (
            f"{company.slug} {obs['source_cell']} {obs['metric']} "
            f"{obs['period_label']}")


def test_every_cim_figure_appears_on_the_page_the_key_cites(company):
    key = _key(company)
    pages = [p.extract_text() or "" for p in PdfReader(
        str(FIXTURES / company.slug / f"{company.slug}_cim_slides.pdf")).pages]
    doc = next(d for d in key["documents"] if d["document_type"] == "cim")
    for obs in doc["observations"]:
        value = obs["value_original"]
        printed = f"({abs(value):,.1f})" if value < 0 else f"{value:,.1f}"
        assert printed in pages[obs["source_page"] - 1], (
            f"{company.slug}: {obs['metric']} {obs['period_label']} "
            f"({printed}) not on p.{obs['source_page']}")


def test_the_cim_is_a_landscape_slide_deck_of_realistic_length(company):
    reader = PdfReader(str(FIXTURES / company.slug /
                           f"{company.slug}_cim_slides.pdf"))
    assert len(reader.pages) >= 80
    box = reader.pages[0].mediabox
    assert (round(float(box.width)), round(float(box.height))) == (960, 540)


def test_the_cap_table_totals_one_hundred_percent(company):
    assert _key(company)["structured_sections"]["cap_table_total_pct"] == \
        pytest.approx(100.0)


def test_unreadable_pages_are_declared(company):
    key = _key(company)
    pages = [p.extract_text() or "" for p in PdfReader(
        str(FIXTURES / company.slug / f"{company.slug}_cim_slides.pdf")).pages]
    declared = key["pages_requiring_ocr_or_flagging"]
    assert declared
    for number in declared:
        assert len(pages[number - 1].strip()) < 200


def test_every_planted_defect_is_declared(company):
    declared = {f["key"] for f in
                _key(company)["expected_reconciliation_findings"]}
    assert declared == {d.key for d in company.defects}


def test_fiscal_periods_have_real_dates_matching_the_year_end(company):
    """A June year end means FY24 runs Jul-23 to Jun-24. Anything
    assuming FY24 is calendar 2024 breaks on Larkspur and Calderwood."""
    for period in company.periods:
        assert period.end.month == company.fiscal_year_end_month
        assert period.start < period.end


# ---------------------------------------------------------------------
# Company-specific defects
# ---------------------------------------------------------------------

def test_larkspur_puts_a_forecast_inside_the_historical_block():
    from tests.fixtures.companies.larkspur import LARKSPUR

    sheet = load_workbook(FIXTURES / "larkspur" /
                          "larkspur_financials.xlsx")["Scorecard"]
    merged = {str(r) for r in sheet.merged_cells.ranges}
    assert "D7:G7" in merged and sheet["D7"].value == "Historicals"

    # Column G is inside the historical merge; its own label says forecast.
    assert sheet.cell(row=8, column=7).value.endswith("E")
    assert LARKSPUR.period("FY25").basis.value == "management_forecast"


def test_larkspur_names_a_rule_of_40_variant_it_cannot_support():
    """Decision 008: a label naming an input the company does not report
    is a hard error, not a warning."""
    from tests.fixtures.companies.larkspur import LARKSPUR

    assert LARKSPUR.rule_of_40_variant.value.startswith("carr_growth")
    metrics = {i.metric for i in LARKSPUR.line_items}
    assert "carr" not in metrics


def test_calderwood_documents_disagree_on_the_same_years_revenue():
    key = _key(next(c for c in COMPANIES if c.slug == "calderwood"))
    values = {}
    for doc in key["documents"]:
        for obs in doc["observations"]:
            if obs["metric"] == "revenue" and obs["period_label"] == "FY24":
                values[doc["document_type"]] = obs["value_standardised"]

    assert values["cim"] == pytest.approx(31_800_000)
    assert values["management_accounts"] == pytest.approx(30_900_000)
    # Both survive. Neither is silently preferred.
    assert len(set(values.values())) == 2


def test_calderwood_presents_opex_on_two_overlapping_axes():
    """Summing every expense row double counts the whole cost base."""
    from storage.enums import BreakdownAxis
    from tests.fixtures.companies.calderwood import CALDERWOOD

    by_nature = [i for i in CALDERWOOD.line_items
                 if i.breakdown_axis == BreakdownAxis.BY_NATURE]
    by_function = [i for i in CALDERWOOD.line_items
                   if i.breakdown_axis == BreakdownAxis.BY_FUNCTION]
    assert by_nature and by_function

    total = CALDERWOOD.item("total_opex").values["FY24"]
    assert sum(i.values["FY24"] for i in by_nature) == pytest.approx(total, abs=1)
    assert sum(i.values["FY24"] for i in by_function) == pytest.approx(total, abs=1)

    naive = sum(i.values["FY24"] for i in by_nature + by_function)
    assert naive == pytest.approx(2 * total, abs=2)


def test_calderwood_carries_a_check_that_cannot_be_evaluated():
    sheet = load_workbook(FIXTURES / "calderwood" /
                          "calderwood_financials.xlsx")["Scorecard"]
    formulas = [sheet.cell(row=r, column=4).value for r in range(20, 45)]
    broken = [f for f in formulas
              if isinstance(f, str) and "Adj_EBITDA_PL" in f]
    assert broken, "the inert check formula is missing"
