"""Does the answer key describe the documents that actually exist?

Generating the key alongside the documents guarantees they came from the
same source. It does not guarantee the generators agree about what they
produced — an off-by-one in a row map, or a page number captured before
a flowable moved, would produce a key that is internally tidy and wrong.

These tests read the rendered files back off disk and check.

If these pass, the fixtures are a yardstick. If they fail, every accuracy
number the project ever reports is measured against a broken ruler.
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from tests.fixtures.companies.thornbury import THORNBURY

FIXTURES = Path(__file__).resolve().parent / "fixtures"
THORNBURY_DIR = FIXTURES / "thornbury"


@pytest.fixture(scope="module")
def key() -> dict:
    path = THORNBURY_DIR / "expected.json"
    if not path.exists():
        pytest.skip("fixtures not built; run python -m tests.fixtures.build_all")
    return json.loads(path.read_text())


@pytest.fixture(scope="module")
def cim_pages() -> list[str]:
    """The slide deck: the primary CIM, matching the shape of a real one."""
    return [p.extract_text() or ""
            for p in PdfReader(str(THORNBURY_DIR / "thornbury_cim_slides.pdf")).pages]


@pytest.fixture(scope="module")
def prose_pages() -> list[str]:
    """The portrait prose CIM: the deliberately easy regression fixture."""
    return [p.extract_text() or ""
            for p in PdfReader(str(THORNBURY_DIR / "thornbury_cim.pdf")).pages]


@pytest.fixture(scope="module")
def sheet():
    return load_workbook(THORNBURY_DIR / "thornbury_financials.xlsx",
                         data_only=False)["Scorecard"]


def _document(key: dict, doc_type: str) -> dict:
    return next(d for d in key["documents"] if d["document_type"] == doc_type)


# ---------------------------------------------------------------------
# The truth itself
# ---------------------------------------------------------------------

def test_the_truth_is_internally_consistent():
    """Before any defect is applied, every total equals its components.

    A fixture built on inconsistent numbers would fail reconciliation for
    reasons unrelated to the planted defects, and we could never tell a
    real finding from an accident.
    """
    assert THORNBURY.check_internal_consistency() == []


# ---------------------------------------------------------------------
# The workbook
# ---------------------------------------------------------------------

def test_every_workbook_figure_is_in_the_cell_the_key_cites(key, sheet):
    observations = _document(key, "financials_workbook")["observations"]
    assert observations, "no workbook observations in the answer key"

    mismatches = []
    for obs in observations:
        actual = sheet[obs["source_cell"]].value
        if actual is None or abs(actual - obs["value_original"]) > 1e-6:
            mismatches.append(
                f"{obs['source_cell']} ({obs['metric']} {obs['period_label']}): "
                f"key says {obs['value_original']}, cell holds {actual}"
            )
    assert not mismatches, "\n".join(mismatches)


def test_the_workbook_states_its_units(sheet):
    """A sheet whose scale is only implied is not safely readable."""
    labels = {sheet.cell(row=r, column=2).value for r in range(1, 10)}
    assert "Units" in labels
    assert sheet["C4"].value == "thousands"


def test_actuals_and_forecasts_are_separated_by_a_merged_header(sheet):
    """Basis is signalled the way real models signal it: a merged cell
    several rows above the numbers, not a suffix on each column."""
    merged = {str(r) for r in sheet.merged_cells.ranges}
    assert "D7:F7" in merged and "G7:I7" in merged
    assert sheet["D7"].value == "Historicals"
    assert sheet["G7"].value == "Management forecast"


# ---------------------------------------------------------------------
# The CIM
# ---------------------------------------------------------------------

def test_the_financial_table_is_on_the_page_the_key_cites(key, cim_pages):
    doc = _document(key, "cim")
    page = doc["section_pages"]["financial_table"]
    text = cim_pages[page - 1]
    assert "Financial Summary" in text
    assert "Turnover" in text and "Gross profit" in text


def test_every_cim_figure_appears_on_the_page_the_key_cites(key, cim_pages):
    """Page numbers are captured during rendering, not guessed. If a
    paragraph above the table changes length, this test catches it."""
    doc = _document(key, "cim")
    missing = []
    for obs in doc["observations"]:
        text = cim_pages[obs["source_page"] - 1]
        value = obs["value_original"]
        printed = f"({abs(value):,.1f})" if value < 0 else f"{value:,.1f}"
        if printed not in text:
            missing.append(f"{obs['metric']} {obs['period_label']}: "
                           f"{printed} not on p.{obs['source_page']}")
    assert not missing, "\n".join(missing)


def test_the_cim_is_a_landscape_slide_deck():
    """A CIM exported from PowerPoint is a different extraction problem
    from a flowing document: text arrives as positioned fragments with no
    guaranteed reading order. Measured on a real IM at 960x540."""
    page = PdfReader(str(THORNBURY_DIR / "thornbury_cim_slides.pdf")).pages[0]
    width, height = float(page.mediabox.width), float(page.mediabox.height)
    assert (round(width), round(height)) == (960, 540)
    assert width / height > 1.55


def test_the_cim_text_arrives_out_of_reading_order(cim_pages):
    """A real deck shows ~18% of text runs moving UP the page. A fixture
    drawn tidily downwards would let us build an extractor that quietly
    depends on reading order and fails on every real document."""
    import statistics

    reader = PdfReader(str(THORNBURY_DIR / "thornbury_cim_slides.pdf"))
    shares = []
    for page in reader.pages:
        ys: list[float] = []

        def visitor(text, cm, tm, font_dict, font_size, _ys=ys):
            if text and text.strip():
                _ys.append(tm[4] * cm[1] + tm[5] * cm[3] + cm[5])

        page.extract_text(visitor_text=visitor)
        if len(ys) > 3:
            back = sum(1 for a, b in zip(ys, ys[1:]) if b > a + 2)
            shares.append(back / (len(ys) - 1))
    assert 0.10 <= statistics.median(shares) <= 0.28


def test_some_pages_carry_almost_no_text(cim_pages):
    """17% of the real deck is dividers and full-page graphics."""
    sparse = [i + 1 for i, p in enumerate(cim_pages) if len(p.strip()) < 200]
    assert 0.10 <= len(sparse) / len(cim_pages) <= 0.28
    assert any(len(cim_pages[i - 1].strip()) == 0 for i in sparse)


def test_unreadable_pages_are_declared_in_the_key(key, cim_pages):
    """Figures that exist only as pixels must be reported as unread.
    Treating them as empty produces a confident answer with a hole in it."""
    declared = key["pages_requiring_ocr_or_flagging"]
    assert declared, "no unreadable pages declared"
    for number in declared:
        assert len(cim_pages[number - 1].strip()) < 200


def test_the_arr_bridge_ties_to_itself_but_not_to_the_accounts(key):
    """The FY24 bridge is internally consistent and disagrees with the
    ARR line — a bridge built before a late adjustment to the accounts."""
    rows = key["structured_sections"]["arr_bridge"]
    assert rows
    for row in rows:
        assert row["closing_implied_by_components"] == pytest.approx(
            row["closing_as_printed"], abs=1.0)

    mismatched = [r["period_label"] for r in rows
                  if abs(r["closing_as_printed"]
                         - r["arr_in_financial_summary"]) > 1.0]
    assert mismatched == ["FY24"]


def test_the_cap_table_totals_one_hundred_percent(key):
    assert key["structured_sections"]["cap_table_total_pct"] == pytest.approx(100.0)


def test_the_deck_carries_the_sections_a_real_cim_has(key):
    sections = key["structured_sections"]
    assert len(sections["team"]) >= 4
    assert len(sections["competitors"]) >= 3
    assert len(sections["retention"]) >= 3


def test_the_prose_fixture_still_builds_as_an_easy_case(prose_pages):
    """Kept for diagnosis: when extraction fails on slides, the same
    prompt against clean prose says whether the problem is the format or
    the logic."""
    assert len(prose_pages) >= 75
    assert any("Financial Summary" in p for p in prose_pages)


def test_the_cim_is_long_enough_to_require_chunking(cim_pages):
    """A fixture small enough to fit in one model call would let us build
    a pipeline that works against a hosted model and fails on the local
    one — which is the only provider confidential deals may use."""
    assert len(cim_pages) >= 80


def test_currency_figures_are_scattered_throughout_the_cim(cim_pages):
    """Retrieval must discriminate, not just find the only page that
    mentions money.

    Measured across pages that carry text at all: a deck's dividers and
    graphic-only slides cannot contain a currency figure, so counting
    them in the denominator would set a threshold no real deck could
    meet.
    """
    content = [p for p in cim_pages if len(p.strip()) >= 200]
    with_money = sum(1 for p in content if "£" in p)
    assert with_money > 0.75 * len(content), (
        f"{with_money} of {len(content)} content pages mention money")


# ---------------------------------------------------------------------
# The planted defects are actually present
# ---------------------------------------------------------------------

def test_the_rekeying_slip_is_present_and_differs_by_exactly_1000x(key, sheet):
    """FY23 cost of sales, re-keyed at the CIM's scale into a sheet
    denominated in thousands."""
    cim = _document(key, "cim")
    wb = _document(key, "financials_workbook")

    cim_cogs = next(o for o in cim["observations"]
                    if o["metric"] == "cogs" and o["period_label"] == "FY23")
    wb_cogs = next(o for o in wb["observations"]
                   if o["metric"] == "cogs" and o["period_label"] == "FY23")

    assert abs(sheet[wb_cogs["source_cell"]].value) == pytest.approx(3.5)
    ratio = abs(cim_cogs["value_standardised"] / wb_cogs["value_standardised"])
    assert ratio == pytest.approx(1000.0)


def test_the_workbook_check_fails_for_fy23_and_no_other_period(key, sheet):
    """Gross profit does not equal revenue less cost of sales — for
    exactly one period."""
    wb = _document(key, "financials_workbook")
    by_period: dict[str, dict[str, float]] = {}
    for obs in wb["observations"]:
        by_period.setdefault(obs["period_label"], {})[obs["metric"]] = \
            obs["value_original"]

    failing = [p for p, v in by_period.items()
               if abs(v["gross_profit"] - (v["revenue"] + v["cogs"])) > 0.5]
    assert failing == ["FY23"]


def test_the_exec_summary_restates_revenue_differently(cim_pages, key):
    """Two figures for FY24 revenue in one document. Both correct: the
    exec summary is continuing operations only."""
    doc = _document(key, "cim")
    exec_text = cim_pages[doc["section_pages"]["executive_summary"] - 1]
    assert "15.1m" in exec_text

    table_text = cim_pages[doc["section_pages"]["financial_table"] - 1]
    assert "15.4" in table_text


def test_the_cim_labels_adjusted_ebitda_as_plain_ebitda(key, cim_pages):
    """Extraction must transcribe the wrong label faithfully. Noticing it
    is reconciliation's job."""
    doc = _document(key, "cim")
    rows = {o["original_label"] for o in doc["observations"]}
    assert "EBITDA" in rows
    assert "Adj. EBITDA" not in rows
    assert "Exceptional items" in rows

    by_period: dict[str, dict[str, float]] = {}
    for obs in doc["observations"]:
        by_period.setdefault(obs["period_label"], {})[obs["metric"]] = \
            obs["value_standardised"]
    fy24 = by_period["FY24"]
    statutory = fy24["gross_profit"] - fy24["total_opex"]
    assert fy24["ebitda"] - statutory == pytest.approx(fy24["ebitda_adjustments"])


def test_every_planted_defect_is_declared_in_the_key(key):
    findings = {f["key"] for f in key["expected_reconciliation_findings"]}
    assert findings == {d.key for d in THORNBURY.defects}
    # One of them is not an error, and must stop being raised once explained.
    assert any(not f["is_genuine_error"]
               for f in key["expected_reconciliation_findings"])


# ---------------------------------------------------------------------
# Reproducibility
# ---------------------------------------------------------------------

def test_rebuilding_produces_an_identical_answer_key(key, tmp_path):
    """The fixtures must be regenerable. If a rebuild moved a page number
    or a cell reference, every stored source reference would rot."""
    from tests.fixtures import build_answer_key, build_cim_slides, build_workbook

    workbook_path, rows = build_workbook.build(THORNBURY, tmp_path)
    _, marks, _ = build_cim_slides.build_to_page_count(THORNBURY, tmp_path, 84)
    key_path = build_answer_key.build(
        THORNBURY, tmp_path, marks, rows,
        cim_filename="thornbury_cim_slides.pdf",
        workbook_filename="thornbury_financials.xlsx",
        unreadable_pages=sorted({marks["unreadable_chart"],
                                 marks["blank_divider"]}))

    assert json.loads(key_path.read_text()) == key
